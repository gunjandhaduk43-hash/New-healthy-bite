import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# --- Color Constants ---
# STRICT RULE 1: NO COLORS IN HEADINGS (Pure Black / 0x000000)
COLOR_HEADING = RGBColor(0x00, 0x00, 0x00)
COLOR_DARK_TEXT = RGBColor(0x00, 0x00, 0x00)
COLOR_MUTED = RGBColor(0x00, 0x00, 0x00)

# STRICT RULE 2: LITE SKY BLUE COLOR IN TABLES
# BAE6FD is clean Tailwind Sky-200 (light sky blue)
COLOR_TABLE_HEADER_HEX = "BAE6FD" 
COLOR_TABLE_HEADER_TEXT = RGBColor(0x00, 0x00, 0x00) # Crisp pure black for maximum contrast

COLOR_BORDER_HEX = "94A3B8"        # Subtle slate border
COLOR_ALT_ROW_HEX = "F0F9FF"       # Very soft sky tint for alternate rows
COLOR_PASS_HEX = "E0F2FE"          # Soft sky blue fill for PASS status
COLOR_PASS_TEXT = RGBColor(0x00, 0x00, 0x00) # Bold pure black text

def set_cell_background(cell, hex_color):
    """Sets cell background shading."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets inner margins for a table cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    """Sets cell borders."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for b_name, b_val in borders.items():
        if b_val:
            node = OxmlElement(f'w:{b_name}')
            node.set(qn('w:val'), b_val.get('val', 'single'))
            node.set(qn('w:sz'), str(b_val.get('sz', 4)))
            node.set(qn('w:space'), '0')
            node.set(qn('w:color'), b_val.get('color', COLOR_BORDER_HEX))
            tcBorders.append(node)
        else:
            node = OxmlElement(f'w:{b_name}')
            node.set(qn('w:val'), 'none')
            tcBorders.append(node)
    tcPr.append(tcBorders)

def add_header_footer(doc):
    """Adds running headers and footers to all sections in pure black."""
    for s in doc.sections:
        s.header_distance = Inches(0.4)
        s.footer_distance = Inches(0.4)
        
        # Header
        header = s.header
        p_head = header.paragraphs[0]
        p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_head = p_head.add_run("Healthy Bite — Digital Restaurant Menu & Food Ordering System")
        r_head.font.name = "Calibri"
        r_head.font.size = Pt(8.5)
        r_head.font.italic = True
        r_head.font.color.rgb = COLOR_HEADING
        
        # Footer
        footer = s.footer
        p_foot = footer.paragraphs[0]
        p_foot.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r_foot_left = p_foot.add_run("Department of Computer Applications • Christ College, Rajkot (Affiliated to Saurashtra University)")
        r_foot_left.font.name = "Calibri"
        r_foot_left.font.size = Pt(8.5)
        r_foot_left.font.color.rgb = COLOR_HEADING

def style_heading_1(p, text):
    """Chapter heading - Strictly NO COLOR (Pure Black)."""
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = COLOR_HEADING
    return p

def style_heading_2(p, text):
    """Section heading - Strictly NO COLOR (Pure Black)."""
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = COLOR_HEADING
    return p

def style_heading_3(p, text):
    """Sub-section heading - Strictly NO COLOR (Pure Black)."""
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = COLOR_HEADING
    return p

def add_body_paragraph(doc, text="", space_after=6, line_spacing=1.15):
    """Standard body paragraph with justified text and pure black color."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if text:
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(10.5)
        run.font.color.rgb = COLOR_HEADING
    return p

def add_bullet_point(doc, bold_prefix, text):
    """Adds a clean bullet point in pure black."""
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Inches(0.25)
    r_bold = p.add_run(bold_prefix)
    r_bold.font.name = "Calibri"
    r_bold.font.size = Pt(10.5)
    r_bold.font.bold = True
    r_bold.font.color.rgb = COLOR_HEADING
    
    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(10.5)
    r_text.font.color.rgb = COLOR_HEADING
    return p

def add_table_styled(doc, headers, data, col_widths=None):
    """
    Renders a table where header cells are shaded with LITE SKY BLUE (#BAE6FD),
    clean pure black text, subtle borders, and alternate row tints.
    """
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    border_spec = {'val': 'single', 'sz': 4, 'color': COLOR_BORDER_HEX}
    
    # 1. Header Row (Lite Sky Blue)
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], COLOR_TABLE_HEADER_HEX)
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        set_cell_borders(hdr_cells[i], top=border_spec, bottom=border_spec, left=border_spec, right=border_spec)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            r.font.bold = True
            r.font.color.rgb = COLOR_TABLE_HEADER_TEXT
            
    # 2. Data Rows
    for row_idx, row_data in enumerate(data):
        row_cells = table.rows[row_idx + 1].cells
        bg_color = COLOR_ALT_ROW_HEX if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(row_data):
            row_cells[col_idx].text = str(val)
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], top=80, bottom=80, left=120, right=120)
            set_cell_borders(row_cells[col_idx], top=border_spec, bottom=border_spec, left=border_spec, right=border_spec)
            p = row_cells[col_idx].paragraphs[0]
            
            # Styling cell content
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(9.5)
                r.font.color.rgb = COLOR_HEADING
            
            # Special alignment logic
            str_val = str(val).strip()
            if col_idx == 0 and ("TC" in str_val or str_val.isdigit() or str_val.startswith("P") or str_val.startswith("Table")):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if p.runs:
                    p.runs[0].font.bold = True
            elif "PASS" in str_val.upper() or "EXECUTED" in str_val.upper():
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_cell_background(row_cells[col_idx], COLOR_PASS_HEX)
                if p.runs:
                    p.runs[0].font.bold = True
                    p.runs[0].font.color.rgb = COLOR_PASS_TEXT
            elif col_idx == 2 and len(headers) == 4 and ("PK" in str_val or "FK" in str_val or "UK" in str_val):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if p.runs:
                    p.runs[0].font.bold = True
                    
    # 3. Apply Column Widths if provided
    if col_widths and len(col_widths) == len(headers):
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = width

    # Add small spacing after table
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(2)
    p_spacer.paragraph_format.space_after = Pt(8)
    return table

def add_image_figure(doc, img_path, caption, width=Inches(5.8)):
    """Inserts a centered image with caption in pure black."""
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=width)
        
        # Caption - Pure black, italic, bold
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run(caption)
        r_cap.font.name = "Calibri"
        r_cap.font.size = Pt(9.5)
        r_cap.font.italic = True
        r_cap.font.bold = True
        r_cap.font.color.rgb = COLOR_HEADING
    else:
        p_err = doc.add_paragraph(f"[Image Missing: {os.path.basename(img_path)}]")
        p_err.paragraph_format.space_after = Pt(8)
        p_err.runs[0].font.color.rgb = COLOR_HEADING

def add_code_snippet(doc, code_str, title=""):
    """
    Inserts a styled code box with a Lite Sky Blue (#BAE6FD) header bar,
    guaranteeing that even code tables use lite sky blue color!
    """
    table = doc.add_table(rows=2, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.rows[0].cells[0].width = Inches(6.5)
    table.rows[1].cells[0].width = Inches(6.5)
    
    border_spec = {'val': 'single', 'sz': 4, 'color': COLOR_BORDER_HEX}
    
    # Header bar: Lite Sky Blue (#BAE6FD)
    hdr_cell = table.rows[0].cells[0]
    hdr_cell.text = title if title else "Source Code Implementation"
    set_cell_background(hdr_cell, COLOR_TABLE_HEADER_HEX)
    set_cell_margins(hdr_cell, top=80, bottom=80, left=140, right=140)
    set_cell_borders(hdr_cell, top=border_spec, bottom=border_spec, left=border_spec, right=border_spec)
    p_hdr = hdr_cell.paragraphs[0]
    if p_hdr.runs:
        p_hdr.runs[0].font.name = "Calibri"
        p_hdr.runs[0].font.size = Pt(9.5)
        p_hdr.runs[0].font.bold = True
        p_hdr.runs[0].font.color.rgb = COLOR_TABLE_HEADER_TEXT
        
    # Code body: Clean White background with pure black Consolas code
    code_cell = table.rows[1].cells[0]
    code_cell.text = ""
    set_cell_background(code_cell, "FFFFFF")
    set_cell_margins(code_cell, top=100, bottom=100, left=140, right=140)
    set_cell_borders(code_cell, top=border_spec, bottom=border_spec, left=border_spec, right=border_spec)
    
    p = code_cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_str)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = COLOR_HEADING
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(6)

print("Updated healthy_bite_doc_formatter.py loaded successfully.")
