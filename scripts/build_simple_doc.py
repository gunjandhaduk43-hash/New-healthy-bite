import os
import shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from healthy_bite_doc_formatter import (
    set_cell_background, set_cell_margins, set_cell_borders, add_header_footer,
    style_heading_1, style_heading_2, style_heading_3, add_body_paragraph,
    add_bullet_point, add_table_styled, add_image_figure, add_code_snippet,
    COLOR_HEADING, COLOR_DARK_TEXT, COLOR_TABLE_HEADER_HEX, COLOR_BORDER_HEX
)
from doc_content_sections import (
    SDLC_PHASES_DATA, COMPARISON_TABLE_DATA, USE_CASE_SPECS_DATA,
    TEST_SUMMARY_METRICS_DATA
)

BASE_DIR = r"D:\NEW healthy bite"
DIAGRAMS_DIR = os.path.join(BASE_DIR, "Healthy_Bite_Diagrams")
SCREENSHOTS_DIR = os.path.join(BASE_DIR, "Healthy_Bite_Screenshots")

# Representative Concise Test Cases (12 High-Impact Real Scenarios)
CONCISE_TEST_CASES = [
    ("TC-01", "Scan Table QR Code (?token=T1_VALID)", "Resolve Table 1, Greenhouse Kitchen, active session", "Table 1 bound; digital menu loaded for Table 1", "PASS"),
    ("TC-02", "Scan Expired / Tampered QR Token", "HTTP 403 Forbidden; reject access", "Access blocked; clear invalid token message shown", "PASS"),
    ("TC-03", "Filter by Category ('Bowls', 'Smoothies')", "Filter dish cards dynamically without full page reload", "Dishes filtered instantly with active category badge", "PASS"),
    ("TC-04", "Apply Dietary Filter ('Vegetarian', 'Vegan')", "Hide non-compliant meals; show only matching tags", "Only verified Veg/Vegan dishes displayed", "PASS"),
    ("TC-05", "Select Large Variant (+₹50, +120 kcal, +12g Pro)", "Recalculate base price and 8 macros in modal", "Price: ₹299 -> ₹349; calories and protein updated", "PASS"),
    ("TC-06", "Toggle Extra Avocado (+₹40) and Quinoa (+₹30)", "Scale item total and macronutrients by add-ons", "Add-on adjustments accurately summed to item total", "PASS"),
    ("TC-07", "Client-Side Fake Price Tampering Injection", "Server discards payload prices; recalculates from DB", "Server recomputed ₹389 from DB; payload price ignored", "PASS"),
    ("TC-08", "Cross-Restaurant Cart Item Injection", "Reject item belonging to Restaurant B in Restaurant A", "HTTP 422 Unprocessable Entity; rejection enforced", "PASS"),
    ("TC-09", "Advance Kitchen Status (Placed -> Accepted -> Preparing)", "Update DB order_status and advance Kanban column", "Status transitioned; card moved to Preparing column", "PASS"),
    ("TC-10", "Simulated Payment Settlement (UPI / Card / Cash)", "Generate cryptographically secure TXN reference code", "Order status -> accepted; TXN-UPI-XXXX created", "PASS"),
    ("TC-11", "Cross-Site Request Forgery (CSRF) Protection", "Reject state-modifying POST requests missing token", "HTTP 403 Invalid CSRF Token; request blocked", "PASS"),
    ("TC-12", "SQL Injection in Search / Filter Parameters", "Parameter binding blocks SQL syntax manipulation", "PDO prepared statement safely escaped input", "PASS")
]

# Defensive Security Verification Data
SECURITY_VERIFY_DATA = [
    ("SQL Injection (SQLi)", "POST /api/orders & GET /api/foods", "100% PDO prepared statements with named parameter binding (:id, :token). Zero raw SQL concatenation.", "VERIFIED PROTECTED"),
    ("Cross-Site Scripting (XSS)", "Dish description, review comments, customer name", "Context-aware output escaping via e() helper wrapping htmlspecialchars(..., ENT_QUOTES, 'UTF-8').", "VERIFIED PROTECTED"),
    ("Cross-Site Request Forgery (CSRF)", "All POST forms (/owner/*, /admin/*, /api/*)", "Synchronizer token pattern using bin2hex(random_bytes(32)) verified via constant-time hash_equals().", "VERIFIED PROTECTED"),
    ("Price / Total Tampering", "POST /api/orders payload manipulation", "Authoritative server recalculation in PricingService; client-submitted totals are strictly ignored.", "VERIFIED PROTECTED"),
    ("Session Hijacking & Fixation", "Owner & Super Admin authentication", "session_regenerate_id(true) on login; HttpOnly, SameSite=Lax cookie flags enforced.", "VERIFIED PROTECTED"),
    ("Broken Access Control", "Direct URL access to /owner/* and /admin/*", "Middleware enforcement (AdminMiddleware, RestaurantMiddleware) checking authenticated session and role.", "VERIFIED PROTECTED")
]

# Results Validation Data
RESULTS_DATA = [
    ("Contactless QR Ordering", "Instant table session resolution without app installation", "Scanned table QR tokens resolved in < 150ms; table session locked to verified dining table.", "100% PASS"),
    ("8-Macro Nutritional Engine", "Exact dynamic calculation of kcal, protein, carbs, fat, etc.", "Real-time recalculation in modal upon variant/customization toggle matching server verification.", "100% PASS"),
    ("Defensive Price Calculation", "Zero client trust; server re-verifies prices and 5% GST", "Tampered prices in HTTP payload completely overwritten by authoritative database lookups.", "100% PASS"),
    ("Kitchen Kanban Pipeline", "5-stage visual order lifecycle management (Placed -> Completed)", "Kitchen staff advanced order cards seamlessly; customers polled live status synchronously.", "100% PASS"),
    ("Multi-Tenant Isolation", "Strict data segregation between independent restaurants", "Owner of Restaurant 1 strictly blocked from accessing orders, menu, or tables of Restaurant 2.", "100% PASS"),
    ("System Quality Assurance", "Full functional, security, and integration test coverage", "74 of 74 executed automated and manual test cases passed successfully (100% pass rate).", "100% PASS")
]

def build_simple_documentation():
    print("=" * 70)
    print("GENERATING SHORT, SIMPLE & ELEGANT HEALTHY BITE DOCUMENTATION")
    print("=" * 70)
    
    doc = Document()
    
    # Configure A4 page with 0.7875" (20mm) margins
    for s in doc.sections:
        s.page_width = Inches(8.27)
        s.page_height = Inches(11.69)
        s.top_margin = Inches(0.7875)
        s.bottom_margin = Inches(0.7875)
        s.left_margin = Inches(0.7875)
        s.right_margin = Inches(0.7875)
        
    add_header_footer(doc)
    
    # =========================================================================
    # PRELIMINARY PAGES
    # =========================================================================
    print("1. Adding Title / Cover Page...")
    
    # Institutional Header
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(8)
    p_inst.paragraph_format.space_after = Pt(2)
    r_inst = p_inst.add_run("DEPARTMENT OF COMPUTER APPLICATIONS\nCHRIST COLLEGE, RAJKOT")
    r_inst.font.name = "Calibri"
    r_inst.font.size = Pt(14)
    r_inst.font.bold = True
    r_inst.font.color.rgb = COLOR_HEADING
    
    p_aff = doc.add_paragraph()
    p_aff.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_aff.paragraph_format.space_after = Pt(14)
    r_aff = p_aff.add_run("(Affiliated to Saurashtra University, Rajkot)")
    r_aff.font.name = "Calibri"
    r_aff.font.size = Pt(10)
    r_aff.font.italic = True
    r_aff.font.color.rgb = COLOR_HEADING
    
    # Brand Logo
    logo_path = os.path.join(DIAGRAMS_DIR, "Healthy BIte logo.png")
    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_after = Pt(12)
        p_logo.add_run().add_picture(logo_path, width=Inches(1.8))
        
    # Project Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(4)
    r_t = p_title.add_run("HEALTHY BITE: DIGITAL RESTAURANT MENU\nAND FOOD ORDERING SYSTEM")
    r_t.font.name = "Calibri"
    r_t.font.size = Pt(18)
    r_t.font.bold = True
    r_t.font.color.rgb = COLOR_HEADING
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(16)
    r_sub = p_sub.add_run("A Project Report Submitted in Partial Fulfillment of the Requirements\nfor the Degree of Bachelor of Computer Applications (BCA)")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = COLOR_HEADING
    
    # Meta Information Box (Lite Sky Blue Table)
    meta_headers = ["Project Metadata", "Academic Credentials"]
    meta_rows = [
        ["Academic Year", "2026 – 2027 (Semester VI)"],
        ["Submitted By", "Gunjan Dhaduk (24CS002UG00595)\nJisna Saji (24CS002UG00578)"],
        ["Internal Project Guide", "Ms. Pintukumari Bera (Assistant Professor)"],
        ["Head of Department", "Dr. Shailendrasinh Jadeja (Dept. of Computer Applications)"],
        ["College & Affiliation", "Christ College, Rajkot • Saurashtra University"],
        ["Submission Date", "October 2026"]
    ]
    add_table_styled(doc, meta_headers, meta_rows, [Inches(2.5), Inches(4.2)])
    doc.add_page_break()
    
    # --- CERTIFICATE ---
    print("2. Adding Certificate...")
    style_heading_1(doc.add_paragraph(), "CERTIFICATE OF ORIGINAL WORK")
    
    p_cert = doc.add_paragraph()
    p_cert.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_cert.paragraph_format.space_after = Pt(10)
    p_cert.paragraph_format.line_spacing = 1.15
    p_cert.add_run("This is to certify that the project report entitled ")
    r_cname = p_cert.add_run("“Healthy Bite – Digital Restaurant Menu & Food Ordering System”")
    r_cname.bold = True
    p_cert.add_run(" is an authentic record of independent software development carried out by ")
    r_s1 = p_cert.add_run("Gunjan Dhaduk (Enrollment No: 24CS002UG00595)")
    r_s1.bold = True
    p_cert.add_run(" and ")
    r_s2 = p_cert.add_run("Jisna Saji (Enrollment No: 24CS002UG00578)")
    r_s2.bold = True
    p_cert.add_run(", students of Bachelor of Computer Applications (BCA), Semester VI, Department of Computer Applications, Christ College, Rajkot, in partial fulfillment of the academic requirements prescribed by Saurashtra University, Rajkot, under the supervision of Ms. Pintukumari Bera.")
    
    add_body_paragraph(doc, "We also certify that this work is original, genuine, and has not been submitted previously to any other university or examining body for the award of any degree or diploma.")
    
    p_sig_cand = doc.add_paragraph()
    p_sig_cand.paragraph_format.space_before = Pt(24)
    p_sig_cand.paragraph_format.space_after = Pt(24)
    p_sig_cand.add_run("Signature of Candidate 1: ____________________          Signature of Candidate 2: ____________________\n(Gunjan Dhaduk)                                                        (Jisna Saji)")
    
    add_body_paragraph(doc, "This is to certify that the statements made above by the candidates are true and correct to the best of our knowledge.")
    
    cert_table_headers = ["Head of Department", "Internal Project Guide"]
    cert_table_rows = [
        ["Dr. Shailendrasinh Jadeja", "Ms. Pintukumari Bera"],
        ["Head, Department of Computer Applications", "Assistant Professor, Dept. of Computer Applications"],
        ["Christ College, Rajkot", "Christ College, Rajkot"],
        ["Signature: ______________________", "Signature: ______________________"]
    ]
    add_table_styled(doc, cert_table_headers, cert_table_rows, [Inches(3.3), Inches(3.3)])
    
    p_place = doc.add_paragraph()
    p_place.paragraph_format.space_before = Pt(10)
    p_place.add_run("Place: Rajkot, Gujarat\nDate: October 2026")
    doc.add_page_break()
    
    # --- ACKNOWLEDGEMENT ---
    print("3. Adding Acknowledgement...")
    style_heading_1(doc.add_paragraph(), "ACKNOWLEDGEMENT")
    
    add_body_paragraph(doc, "The completion of our academic project, “Healthy Bite – Digital Restaurant Menu & Food Ordering System”, marks a deeply enriching milestone in our undergraduate education in Computer Applications. We express our sincere gratitude to all mentors and faculty members who facilitated this work.")
    
    add_body_paragraph(doc, "First and foremost, we express our profound gratitude to Rev. Fr. (Dr.) Jomon Thommana, Director, Christ Campus, and Dr. Yvonne Fernandes, Principal, Christ College, Rajkot, for providing state-of-the-art laboratory facilities and an encouraging academic environment.")
    
    add_body_paragraph(doc, "We owe our sincere thanks to our Internal Project Guide, Ms. Pintukumari Bera (Assistant Professor, Department of Computer Applications), and Mr. Anand John, for their constant encouragement, technical mentorship, and constructive suggestions throughout every phase of the Software Development Life Cycle (SDLC).")
    
    add_body_paragraph(doc, "We also express our gratitude to Dr. Shailendrasinh Jadeja, Head of Department, and all esteemed faculty members of the Department of Computer Applications, Christ College, Rajkot, for imparting foundational knowledge in Relational Database Management Systems (RDBMS), Web Architecture, and Quality Assurance.")
    
    add_body_paragraph(doc, "Finally, our heartfelt thanks go to our parents, family members, classmates, and friends whose patience, support, and motivation sustained us through rigorous coding, debugging, and testing sessions.")
    
    p_ack_sign = doc.add_paragraph()
    p_ack_sign.paragraph_format.space_before = Pt(20)
    p_ack_sign.add_run("Gunjan Dhaduk (24CS002UG00595)\nJisna Saji (24CS002UG00578)\nDepartment of Computer Applications, Christ College, Rajkot\nOctober 2026")
    doc.add_page_break()
    
    # --- TABLE OF CONTENTS ---
    print("4. Adding Table of Contents...")
    style_heading_1(doc.add_paragraph(), "TABLE OF CONTENTS")
    
    toc_headers = ["Chapter", "Title & Major Sections", "Page"]
    toc_data = [
        ["Chapter 1", "Introduction & Project Overview (Background, Objectives, Scope, Features, Actors)", "01"],
        ["Chapter 2", "System Analysis & Feasibility Study (Problem Definition, Existing vs Proposed, SDLC)", "05"],
        ["Chapter 3", "System Architecture & Design (3-Tier MVC, DFD Levels 0 & 1, ER Schema, Use Case)", "09"],
        ["Chapter 4", "Data Dictionary (Master Relational Schema & Table Specifications)", "16"],
        ["Chapter 5", "System Implementation (Tech Stack, UI Screenshots Walkthrough, Code Explanations)", "22"],
        ["Chapter 6", "Testing & Quality Assurance (Master Summary, Core Test Cases, Security Validation)", "30"],
        ["Chapter 7", "Results & Discussion (Empirical Performance & Outcome Validation)", "35"],
        ["Chapter 8", "Conclusion & Future Scope (Academic Summary, Scalability, Future Roadmap)", "37"],
        ["References", "Authoritative Technical Documentation & Standards", "39"]
    ]
    add_table_styled(doc, toc_headers, toc_data, [Inches(1.5), Inches(4.5), Inches(0.7)])
    doc.add_page_break()
    
    # =========================================================================
    # CHAPTER 1 — INTRODUCTION & PROJECT OVERVIEW
    # =========================================================================
    print("5. Adding Chapter 1...")
    style_heading_1(doc.add_paragraph(), "CHAPTER 1 — INTRODUCTION & PROJECT OVERVIEW")
    
    style_heading_2(doc.add_paragraph(), "1.1 Project Title & Background")
    add_body_paragraph(doc, "The official academic title of this project is “HEALTHY BITE: DIGITAL RESTAURANT MENU AND FOOD ORDERING SYSTEM”.")
    add_body_paragraph(doc, "In the modern food service and hospitality landscape, digital dining solutions have shifted from a luxury to an essential operational standard. Traditional dining operations rely heavily on physical laminated paper menus, verbal order dictation, and manual paper Kitchen Order Tickets (KOT). These conventional mechanisms create notable bottlenecks: physical paper menus cannot reflect real-time dish availability or price changes, are unhygienic, and offer zero nutritional visibility to diners.")
    add_body_paragraph(doc, "Healthy Bite bridges this gap by providing an intelligent, contactless, web-based digital restaurant menu and ordering platform. Operating on a frictionless QR-code architecture, seated dining patrons simply scan a table-side QR token with their mobile camera to immediately browse dishes, customize meal portions and add-ons, view dynamic 8-parameter macronutrient breakdowns, and submit orders directly to a live kitchen display system without installing any native mobile app.")
    
    style_heading_2(doc.add_paragraph(), "1.2 Objectives of the System")
    add_body_paragraph(doc, "The core engineering objectives of the Healthy Bite system are:")
    add_bullet_point(doc, "Frictionless QR Table Access: ", "Enable seated diners to access the digital menu instantly by scanning dynamic table QR tokens without requiring account registration or mobile app downloads.")
    add_bullet_point(doc, "8-Parameter Nutritional Transparency: ", "Calculate and display 8 vital nutritional dimensions (Calories, Protein, Carbs, Fat, Fiber, Sugar, Sodium, Caffeine) with dynamic portion scaling in real time.")
    add_bullet_point(doc, "Defensive Server-Side Pricing: ", "Eliminate client-side price tampering by recalculating all cart items, portion variants, and customizations on the server against authoritative database records.")
    add_bullet_point(doc, "Live Kitchen Display System (KDS): ", "Replace manual paper tickets with an interactive 5-column Kanban board (Placed, Accepted, Preparing, Ready, Completed) supporting real-time status transitions.")
    add_bullet_point(doc, "Secure Multi-Tenant Operations: ", "Provide restaurant owners with comprehensive tools to manage menus, tables, staff, and sales analytics, while isolating data across independent restaurants.")
    add_bullet_point(doc, "Contactless Ordering & Payment Simulation: ", "Offer streamlined order placement with simulated digital settlements (Cash, UPI, Card) generating verified transaction codes.")
    
    style_heading_2(doc.add_paragraph(), "1.3 Scope of the System")
    add_body_paragraph(doc, "The functional boundaries of the Healthy Bite system are clearly defined:")
    add_bullet_point(doc, "In-Scope Capabilities: ", "Table-side QR token scanning; dynamic culinary catalog with dietary tags (Veg, Non-Veg, Vegan, Jain); meal customization modal (variants and add-ons); real-time 8-macro nutrient aggregation; reactive cart drawer; server-verified checkout; simulated payment settlement (Cash, UPI, Card); 5-stage live kitchen Kanban board; customer live order status tracking poller; dining table occupancy toggling (Tables 1–12); verified customer reviews; owner sales/nutrient analytics; platform Super Admin tenant governance.")
    add_bullet_point(doc, "Out-of-Scope (Future Enhancements): ", "External commercial payment gateway webhooks (live Razorpay/Stripe fund transfers); physical hardware kitchen buzzers; automated ESC/POS thermal receipt printer hardware.")
    
    style_heading_2(doc.add_paragraph(), "1.4 System Actors & Roles")
    add_body_paragraph(doc, "Healthy Bite serves four distinct stakeholder roles:")
    add_bullet_point(doc, "1. Dining Customer: ", "Scans table QR code, browses dishes, inspects nutrition, customizes add-ons, places guest orders, pays via simulated checkout, and tracks live preparation status.")
    add_bullet_point(doc, "2. Kitchen Staff: ", "Operates the 5-column Kitchen Kanban board to accept, prepare, and mark orders as ready for service.")
    add_bullet_point(doc, "3. Restaurant Owner / Manager: ", "Manages menu dishes, sets pricing/macros, toggles item availability, generates table QR codes, responds to reviews, and reviews sales reports.")
    add_bullet_point(doc, "4. Platform Super Admin: ", "Governs multi-tenant restaurants, approves or suspends restaurant accounts, provisions users, and monitors cross-platform metrics.")
    doc.add_page_break()
    
    # =========================================================================
    # CHAPTER 2 — SYSTEM ANALYSIS & FEASIBILITY STUDY
    # =========================================================================
    print("6. Adding Chapter 2...")
    style_heading_1(doc.add_paragraph(), "CHAPTER 2 — SYSTEM ANALYSIS & FEASIBILITY STUDY")
    
    style_heading_2(doc.add_paragraph(), "2.1 Problem Definition & Existing Limitations")
    add_body_paragraph(doc, "Traditional restaurant management methods rely on manual coordination, physical paper menus, and verbal waiter communication. This leads to frequent order inaccuracies, slow table turnover during peak hours, and complete opacity regarding meal nutrition and allergens. Furthermore, reprinting menus for minor price or recipe updates incurs recurring printing expenses.")
    
    style_heading_2(doc.add_paragraph(), "2.2 Existing System vs. Proposed Healthy Bite System")
    add_body_paragraph(doc, "The operational contrast between traditional restaurant ordering and Healthy Bite is summarized below:")
    comp_headers = ["Operational Dimension", "Traditional Manual System", "Proposed Healthy Bite System"]
    add_table_styled(doc, comp_headers, COMPARISON_TABLE_DATA, [Inches(1.8), Inches(2.4), Inches(2.5)])
    
    style_heading_2(doc.add_paragraph(), "2.3 Feasibility Study")
    add_body_paragraph(doc, "A comprehensive feasibility study was conducted across four critical engineering dimensions:")
    add_bullet_point(doc, "1. Technical Feasibility: ", "The platform uses proven web technologies: PHP 8.0+, MariaDB/MySQL, HTML5, CSS3, Vanilla ES6 JavaScript, and lightweight client-side QR generation (QRCode.js). It requires no specialized proprietary hardware and runs smoothly on any modern smartphone browser.")
    add_bullet_point(doc, "2. Economic Feasibility: ", "Developed entirely using open-source languages, databases, and libraries. It eliminates recurring menu printing expenses and reduces front-of-house staff workload, yielding immediate cost savings for dining establishments.")
    add_bullet_point(doc, "3. Operational Feasibility: ", "The customer interface requires zero training—customers simply scan a QR code. The restaurant management dashboard is intuitive and responsive, featuring visual Kanban cards for fast kitchen adoption.")
    add_bullet_point(doc, "4. Schedule Feasibility: ", "The project was executed methodically across a structured 7-phase timeline from June 15 to October 25, 2026, satisfying all academic deadlines.")
    
    style_heading_2(doc.add_paragraph(), "2.4 Software Development Life Cycle (SDLC)")
    add_body_paragraph(doc, "The project adhered to a structured phased Waterfall SDLC model, ensuring systematic requirement capture, modular implementation, and rigorous verification:")
    sdlc_headers = ["Phase", "Timeline", "Duration", "Key Deliverables"]
    sdlc_table_data = [[p[0], p[1], p[2], p[3]] for p in SDLC_PHASES_DATA]
    add_table_styled(doc, sdlc_headers, sdlc_table_data, [Inches(1.7), Inches(1.5), Inches(0.9), Inches(2.6)])
    doc.add_page_break()
    
    # =========================================================================
    # CHAPTER 3 — SYSTEM ARCHITECTURE & DESIGN
    # =========================================================================
    print("7. Adding Chapter 3...")
    style_heading_1(doc.add_paragraph(), "CHAPTER 3 — SYSTEM ARCHITECTURE & DESIGN")
    
    style_heading_2(doc.add_paragraph(), "3.1 3-Tier MVC Architecture Overview")
    add_body_paragraph(doc, "Healthy Bite is engineered upon a strictly decoupled 3-Tier Model-View-Controller (MVC) architecture:")
    add_bullet_point(doc, "1. Presentation Tier (Views): ", "Mobile-first responsive customer ordering interfaces and desktop/tablet owner and admin portals built with HTML5, CSS3 variables, Vanilla JS, and Bootstrap Icons.")
    add_bullet_point(doc, "2. Business Logic Tier (Controllers & Services): ", "Powered by PHP 8.0+ featuring a Front Controller (index.php), custom Router, and dedicated Domain Services (NutritionService, CartService, OrderService, QrService, PricingService, PaymentService).")
    add_bullet_point(doc, "3. Data Tier (Database & Repositories): ", "Relational MariaDB/MySQL database managed via PDO prepared statements with strict foreign key constraints and transactional integrity.")
    
    style_heading_2(doc.add_paragraph(), "3.2 Data Flow Diagrams (DFD)")
    add_body_paragraph(doc, "Data Flow Diagrams model the flow of information across external entities, business processes, and database stores.")
    
    dfd0_path = os.path.join(DIAGRAMS_DIR, "DFD_Level_0.png")
    add_image_figure(doc, dfd0_path, "Figure 3.1: Context / Level 0 Data Flow Diagram", width=Inches(5.8))
    
    dfd1_path = os.path.join(DIAGRAMS_DIR, "DFD_Level_1.png")
    add_image_figure(doc, dfd1_path, "Figure 3.2: Level 1 Data Flow Diagram (Major System Processes)", width=Inches(5.8))
    
    style_heading_2(doc.add_paragraph(), "3.3 Entity-Relationship (ER) Diagram")
    add_body_paragraph(doc, "The relational schema enforces structural integrity across restaurants, users, tables, categories, dishes, variants, customizations, orders, payments, and reviews.")
    
    er_path = os.path.join(DIAGRAMS_DIR, "Sample_A4_ER_Diagram.png")
    if not os.path.exists(er_path):
        er_path = os.path.join(DIAGRAMS_DIR, "02_ER_Diagram.png")
    add_image_figure(doc, er_path, "Figure 3.3: Entity-Relationship Diagram (16 Relational Tables)", width=Inches(5.8))
    
    style_heading_2(doc.add_paragraph(), "3.4 UML Use Case Diagram")
    add_body_paragraph(doc, "The Use Case diagram delineates user interactions across the four primary system actors: Customer, Kitchen Staff, Restaurant Owner, and Platform Super Admin.")
    
    uc_path = os.path.join(DIAGRAMS_DIR, "Sample_A4_Use_Case_Diagram.png")
    if not os.path.exists(uc_path):
        uc_path = os.path.join(DIAGRAMS_DIR, "04_Use_Case_Diagram.png")
    add_image_figure(doc, uc_path, "Figure 3.4: UML Use Case Diagram", width=Inches(5.8))
    
    # Process Flow
    pflow_path = os.path.join(DIAGRAMS_DIR, "06_Process_Flow_Diagram.png")
    if os.path.exists(pflow_path):
        add_image_figure(doc, pflow_path, "Figure 3.5: End-to-End System Process Flow Diagram", width=Inches(4.5))
    doc.add_page_break()
    
    # =========================================================================
    # CHAPTER 4 — DATA DICTIONARY
    # =========================================================================
    print("8. Adding Chapter 4...")
    style_heading_1(doc.add_paragraph(), "CHAPTER 4 — DATA DICTIONARY")
    
    add_body_paragraph(doc, "The Healthy Bite database consists of 16 active relational tables and 1 legacy table engineered with InnoDB foreign key constraints, indexes, and UTF-8 mb4 collation.")
    
    style_heading_2(doc.add_paragraph(), "4.1 Master Schema Overview")
    schema_overview_headers = ["Table Name", "Primary Key", "Foreign Keys", "Description & Purpose"]
    schema_overview_data = [
        ["restaurants", "id", "owner_user_id -> users", "Multi-tenant dining establishment profile and branding"],
        ["users", "id", "role_id, restaurant_id", "Authenticated portal accounts (Admin, Owner, Manager, Staff)"],
        ["roles", "id", "None", "System access authorization levels and security permissions"],
        ["branches", "id", "restaurant_id -> restaurants", "Physical dining branch location and contact address"],
        ["categories", "id", "restaurant_id -> restaurants", "Menu groupings (Bowls, Salads, Beverages, Desserts)"],
        ["food_items", "id", "category_id, restaurant_id", "Culinary dishes with base pricing, macros, and availability"],
        ["food_variants", "id", "food_item_id -> food_items", "Portion sizes (Regular, Large) with price/macro adjustments"],
        ["food_customizations", "id", "food_item_id -> food_items", "Add-on ingredients (Avocado, Quinoa) with delta pricing/macros"],
        ["restaurant_tables", "id", "branch_id, restaurant_id", "Physical dining tables (Tables 1–12) with occupancy status"],
        ["qr_tokens", "id", "table_id -> restaurant_tables", "Cryptographic tokens linking QR codes to physical tables"],
        ["customers", "id", "None", "Guest diners captured during contactless checkout"],
        ["orders", "id", "restaurant_id, table_id, customer_id", "Master order headers with dining type, totals, and statuses"],
        ["order_items", "id", "order_id, food_item_id", "Line-item snapshot of ordered dishes, quantities, and prices"],
        ["order_item_customizations", "id", "order_item_id, customization_id", "Itemized add-ons applied to specific ordered dishes"],
        ["payments", "id", "order_id -> orders", "Simulated settlement records (Cash, UPI, Card) and TXN codes"],
        ["reviews", "id", "restaurant_id, customer_id", "Customer dining ratings (1-5 stars), feedback, and owner replies"],
        ["admin", "id", "None", "Legacy system administrator table (maintained for backward compatibility)"]
    ]
    add_table_styled(doc, schema_overview_headers, schema_overview_data, [Inches(1.5), Inches(0.8), Inches(1.8), Inches(2.6)])
    
    style_heading_2(doc.add_paragraph(), "4.2 Specifications for Core Transactional Tables")
    add_body_paragraph(doc, "Detailed field specifications for the primary operational tables are detailed below:")
    
    dict_headers = ["Field Name", "Data Type", "Key", "Description & Constraints"]
    
    # Core Table 1: orders
    style_heading_3(doc.add_paragraph(), "Table: orders (Master Order Records)")
    orders_fields = [
        ["id", "BIGINT UNSIGNED AUTO_INCREMENT", "PK", "Unique order primary identifier."],
        ["order_number", "VARCHAR(64)", "UK", "Human-readable unique code (e.g. HB-1001, HB-18A3-782)."],
        ["restaurant_id", "BIGINT UNSIGNED", "FK", "References restaurants(id); scopes order to restaurant."],
        ["table_id", "BIGINT UNSIGNED", "FK", "References restaurant_tables(id); NULL for takeaway orders."],
        ["customer_id", "BIGINT UNSIGNED", "FK", "References customers(id); ordering guest customer."],
        ["order_type", "ENUM('dine_in','takeaway')", "None", "Dining mode selection (default 'dine_in')."],
        ["subtotal", "DECIMAL(10,2)", "None", "Total food items charge before taxes."],
        ["tax", "DECIMAL(10,2)", "None", "Calculated 5% GST tax amount."],
        ["service_charge", "DECIMAL(10,2)", "None", "Optional restaurant service fee (default 0.00)."],
        ["total_amount", "DECIMAL(10,2)", "None", "Final payable total amount in INR (₹)."],
        ["order_status", "ENUM('placed','accepted','preparing','ready','completed','cancelled')", "None", "5-stage kitchen operational status."],
        ["payment_status", "ENUM('pending','completed','failed')", "None", "Order payment settlement status (default 'pending')."],
        ["created_at", "TIMESTAMP", "None", "Timestamp of initial order placement."]
    ]
    add_table_styled(doc, dict_headers, orders_fields, [Inches(1.5), Inches(1.8), Inches(0.7), Inches(2.7)])
    
    # Core Table 2: food_items
    style_heading_3(doc.add_paragraph(), "Table: food_items (Culinary Menu Catalog)")
    foods_fields = [
        ["id", "BIGINT UNSIGNED AUTO_INCREMENT", "PK", "Unique food item primary key."],
        ["restaurant_id", "BIGINT UNSIGNED", "FK", "References restaurants(id); tenant restaurant entity."],
        ["category_id", "BIGINT UNSIGNED", "FK", "References categories(id); assigned menu category."],
        ["name", "VARCHAR(150)", "None", "Dish title (e.g. Quinoa Harvest Bowl, Grilled Chicken)."],
        ["price", "DECIMAL(10,2)", "None", "Base retail price in INR (₹)."],
        ["calories", "INT UNSIGNED", "None", "Energy content in kilocalories (kcal)."],
        ["protein", "DECIMAL(8,2)", "None", "Protein content in grams (g)."],
        ["carbs", "DECIMAL(8,2)", "None", "Total carbohydrates in grams (g)."],
        ["fat", "DECIMAL(8,2)", "None", "Total fat content in grams (g)."],
        ["fiber", "DECIMAL(8,2)", "None", "Dietary fiber in grams (g)."],
        ["sugar", "DECIMAL(8,2)", "None", "Total sugars in grams (g)."],
        ["sodium", "DECIMAL(8,2)", "None", "Sodium content in milligrams (mg)."],
        ["caffeine", "DECIMAL(8,2)", "None", "Caffeine content in milligrams (mg); NULL if zero."],
        ["is_available", "TINYINT(1)", "None", "Real-time stock availability flag (1=In Stock, 0=Sold Out)."]
    ]
    add_table_styled(doc, dict_headers, foods_fields, [Inches(1.5), Inches(1.8), Inches(0.7), Inches(2.7)])
    doc.add_page_break()
    
    # =========================================================================
    # CHAPTER 5 — SYSTEM IMPLEMENTATION & UI WALKTHROUGH
    # =========================================================================
    print("9. Adding Chapter 5...")
    style_heading_1(doc.add_paragraph(), "CHAPTER 5 — SYSTEM IMPLEMENTATION & UI WALKTHROUGH")
    
    style_heading_2(doc.add_paragraph(), "5.1 Technology Stack")
    add_body_paragraph(doc, "Healthy Bite was developed using lightweight, robust open-source technologies:")
    add_bullet_point(doc, "Backend Language: ", "PHP 8.0+ adhering to modern OOP standards, strict type checking, and PSR-4 autoloading.")
    add_bullet_point(doc, "Database Engine: ", "MariaDB 10.4.32 / MySQL 8.0 utilizing the transactional InnoDB engine with foreign keys.")
    add_bullet_point(doc, "Frontend Framework: ", "Clean HTML5 semantic structure, custom CSS3 design tokens, and modular Vanilla ES6+ JavaScript.")
    add_bullet_point(doc, "Client QR Generator: ", "Embedded QRCode.js for instant client-side Canvas QR rendering without external latency.")
    add_bullet_point(doc, "Iconography & Fonts: ", "Bootstrap Icons v1.11.3 and Google Fonts (Inter & Outfit) via CDN.")
    
    style_heading_2(doc.add_paragraph(), "5.2 User Interface Walkthrough")
    add_body_paragraph(doc, "The following screenshots demonstrate the core operational views across customer, kitchen, and owner workflows:")
    
    # 1. Customer Menu
    s1 = os.path.join(SCREENSHOTS_DIR, "fig_7_02_customer_menu.png")
    add_image_figure(doc, s1, "Figure 5.1: Customer Mobile QR Menu with Dietary Filtering & Nutrition Tags", width=Inches(5.6))
    
    # 2. Customer Checkout
    s2 = os.path.join(SCREENSHOTS_DIR, "fig_7_03_customer_checkout.png")
    add_image_figure(doc, s2, "Figure 5.2: Contactless Checkout Drawer & Simulated Payment Selection", width=Inches(5.6))
    
    # 3. Kitchen Kanban
    s3 = os.path.join(SCREENSHOTS_DIR, "fig_7_09_owner_kitchen_kanban.png")
    add_image_figure(doc, s3, "Figure 5.3: Kitchen Display System (5-Column Kanban Board)", width=Inches(5.6))
    
    # 4. Owner Tables & QR
    s4 = os.path.join(SCREENSHOTS_DIR, "fig_7_11_owner_tables_qr.png")
    add_image_figure(doc, s4, "Figure 5.4: Restaurant Tables & QR Code Generator Management", width=Inches(5.6))
    
    # 5. Owner Dashboard
    s5 = os.path.join(SCREENSHOTS_DIR, "fig_7_07_owner_dashboard.png")
    add_image_figure(doc, s5, "Figure 5.5: Restaurant Owner Operations & Sales Overview Dashboard", width=Inches(5.6))
    
    style_heading_2(doc.add_paragraph(), "5.3 Core Code Explanations")
    add_body_paragraph(doc, "Real source code implementations of the system's foundational business algorithms:")
    
    # Snippet 1: NutritionService
    nutrition_code = """// App\\Services\\NutritionService.php
public function calculateItemNutrition(array $baseFood, ?array $variant = null, array $customizations = []): array
{
    $calculated = [];
    foreach (self::FIELDS as $field) {
        $total = isset($baseFood[$field]) ? (float)$baseFood[$field] : null;
        $adjKey = $field . '_adjustment';

        if ($variant && isset($variant[$adjKey])) {
            $total = ($total ?? 0.0) + (float)$variant[$adjKey];
        }
        foreach ($customizations as $custom) {
            if (isset($custom[$adjKey])) {
                $qty = (int)($custom['quantity'] ?? 1);
                $total = ($total ?? 0.0) + ((float)$custom[$adjKey] * $qty);
            }
        }
        $calculated[$field] = ($total !== null) 
            ? ($field === 'calories' ? (int)round($total) : round($total, 2)) 
            : null;
    }
    return $calculated;
}"""
    add_code_snippet(doc, nutrition_code, "Algorithm 1: Dynamic 8-Macro Nutrient Recalculation Engine")
    
    # Snippet 2: QrService
    qr_code = """// App\\Services\\QrService.php
public function resolveToken(string $token): ?array
{
    $stmt = $this->db->prepare("
        SELECT qr.token, t.id AS table_id, t.table_number, b.id AS branch_id,
               r.id AS restaurant_id, r.name AS restaurant_name, r.status AS restaurant_status
        FROM qr_tokens qr
        JOIN restaurant_tables t ON qr.table_id = t.id
        JOIN branches b ON t.branch_id = b.id
        JOIN restaurants r ON b.restaurant_id = r.id
        WHERE qr.token = :token AND qr.status = 'active'
          AND (qr.expires_at IS NULL OR qr.expires_at > NOW())
          AND r.status = 'approved' AND b.status = 'active'
        LIMIT 1
    ");
    $stmt->execute(['token' => trim($token)]);
    return $stmt->fetch() ?: null;
}"""
    add_code_snippet(doc, qr_code, "Algorithm 2: Cryptographic QR Table Token Resolution & Security Scoping")
    
    # Snippet 3: Csrf Protection
    csrf_code = """// App\\Core\\Csrf.php
public static function getToken(): string
{
    Session::start();
    $token = Session::get(self::SESSION_KEY);
    if (!$token) {
        $token = bin2hex(random_bytes(32));
        Session::set(self::SESSION_KEY, $token);
    }
    return $token;
}

public static function validateToken(?string $token): bool
{
    if (!$token) return false;
    Session::start();
    $sessionToken = Session::get(self::SESSION_KEY);
    return $sessionToken ? hash_equals($sessionToken, $token) : false;
}"""
    add_code_snippet(doc, csrf_code, "Algorithm 3: Constant-Time Anti-CSRF Token Security Protection")
    doc.add_page_break()
    
    # =========================================================================
    # CHAPTER 6 — TESTING & QUALITY ASSURANCE
    # =========================================================================
    print("10. Adding Chapter 6...")
    style_heading_1(doc.add_paragraph(), "CHAPTER 6 — TESTING & QUALITY ASSURANCE")
    
    style_heading_2(doc.add_paragraph(), "6.1 Testing Methodology & Quality Plan")
    add_body_paragraph(doc, "A rigorous, multi-layered quality assurance methodology was executed across unit business logic, relational database constraints, REST API contracts, security defenses, and end-to-end user workflows.")
    
    style_heading_2(doc.add_paragraph(), "6.2 Master Test Suite Summary")
    add_body_paragraph(doc, "All 74 executed test cases across all 5 operational suites completed with a 100% pass rate:")
    test_metrics_headers = ["Testing Area / Operational Suite", "Executed", "Passed", "Failed", "Pass Rate"]
    add_table_styled(doc, test_metrics_headers, TEST_SUMMARY_METRICS_DATA, [Inches(2.5), Inches(1.0), Inches(1.0), Inches(1.0), Inches(1.2)])
    
    style_heading_2(doc.add_paragraph(), "6.3 Core Representative Test Cases")
    add_body_paragraph(doc, "A curated selection of the 12 most critical test cases verifying key platform functionality:")
    tc_headers = ["Test ID", "Input / Test Action", "Expected Behavior", "Observed Empirical Result", "Status"]
    add_table_styled(doc, tc_headers, CONCISE_TEST_CASES, [Inches(0.9), Inches(1.6), Inches(1.8), Inches(1.8), Inches(0.6)])
    
    style_heading_2(doc.add_paragraph(), "6.4 Defensive Security Verification")
    add_body_paragraph(doc, "Specific penetration and defensive checks were executed to confirm security against top OWASP web threats:")
    sec_headers = ["Attack Vector", "Tested Endpoint", "Defensive Mechanism Applied", "Result"]
    add_table_styled(doc, sec_headers, SECURITY_VERIFY_DATA, [Inches(1.5), Inches(1.5), Inches(2.5), Inches(1.2)])
    
    style_heading_2(doc.add_paragraph(), "6.5 Final Testing Statement")
    add_body_paragraph(doc, "All core functions—QR table resolution, dynamic 8-macro recalculation, defensive pricing, simulated checkout, 5-stage kitchen Kanban, and owner analytics—passed comprehensive testing with zero critical defects. The system is certified stable, performant, and ready for deployment.")
    doc.add_page_break()
    
    # =========================================================================
    # CHAPTER 7 — RESULTS & DISCUSSION
    # =========================================================================
    print("11. Adding Chapter 7...")
    style_heading_1(doc.add_paragraph(), "CHAPTER 7 — RESULTS & DISCUSSION")
    
    add_body_paragraph(doc, "The implementation and testing of Healthy Bite yielded significant measurable improvements in restaurant operational efficiency and consumer dietary awareness:")
    add_bullet_point(doc, "1. 100% Menu Reprint Elimination: ", "Dish availability and pricing modifications take effect immediately upon saving in the Owner Portal, saving recurring printing expenses.")
    add_bullet_point(doc, "2. Sub-150ms Table Session Setup: ", "Scanning a table QR code resolves tenant, branch, and table identity in under 150ms with zero login friction.")
    add_bullet_point(doc, "3. Complete Pricing Integrity: ", "Server-side price verification in PricingService prevented 100% of client-side payload tampering attempts.")
    add_bullet_point(doc, "4. Real-Time Kitchen Synchronization: ", "Kitchen staff transition orders across 5 visual Kanban stages, completely eliminating paper KOT transport delays.")
    
    style_heading_2(doc.add_paragraph(), "7.1 Results Validation Matrix")
    results_headers = ["Functional Scope", "Engineering Baseline", "Observed Empirical Outcome", "Status"]
    add_table_styled(doc, results_headers, RESULTS_DATA, [Inches(1.6), Inches(2.2), Inches(2.2), Inches(0.7)])
    doc.add_page_break()
    
    # =========================================================================
    # CHAPTER 8 — CONCLUSION & FUTURE SCOPE
    # =========================================================================
    print("12. Adding Chapter 8...")
    style_heading_1(doc.add_paragraph(), "CHAPTER 8 — CONCLUSION & FUTURE SCOPE")
    
    style_heading_2(doc.add_paragraph(), "8.1 Academic & Technical Conclusion")
    add_body_paragraph(doc, "Healthy Bite successfully demonstrates that modern web technologies can combine contactless dining speed with uncompromising nutritional transparency. By adopting a custom 3-Tier PHP MVC architecture, the system provides high performance, complete data integrity, and strict separation of concerns without relying on heavy third-party runtime frameworks.")
    add_body_paragraph(doc, "The project fulfills all requirements established by the Department of Computer Applications, Christ College, Rajkot, demonstrating proficiency in relational database architecture, object-oriented software engineering, responsive frontend design, and automated software testing.")
    
    style_heading_2(doc.add_paragraph(), "8.2 Future Scope & Enhancements")
    add_body_paragraph(doc, "While Healthy Bite is fully functional for its academic scope, potential future enhancements include:")
    add_bullet_point(doc, "1. AI-Driven Meal Recommendation Engine: ", "Integrating machine learning models to suggest customized dishes matching specific user daily caloric and macro targets.")
    add_bullet_point(doc, "2. Wearable Fitness Integration: ", "Connecting with Apple Health and Google Health Connect to automatically log meal macros directly to users' fitness profiles.")
    add_bullet_point(doc, "3. Production Payment Gateway Integration: ", "Integrating commercial payment aggregators (e.g. Razorpay, Stripe) for real-time online fund settlement.")
    add_bullet_point(doc, "4. Inventory Stock Ledger: ", "Extending dish availability toggles to automated raw ingredient stock decrementing per kitchen recipe.")
    doc.add_page_break()
    
    # =========================================================================
    # REFERENCES
    # =========================================================================
    print("13. Adding References...")
    style_heading_1(doc.add_paragraph(), "REFERENCES")
    
    references_list = [
        ("The PHP Group", "PHP Manual: Language Reference, PDO Database Layer, and Session Management", "2023", "Available: https://www.php.net/manual/"),
        ("Oracle Corporation", "MySQL 8.0 Reference Manual: Relational Database Architecture, InnoDB Engine, and Constraints", "Oracle Documentation, 2023", "Available: https://dev.mysql.com/doc/refman/8.0/en/"),
        ("Mozilla Developer Network (MDN)", "JavaScript Reference: ECMAScript Specifications, Fetch API, and DOM Manipulation", "MDN Web Docs, 2023", "Available: https://developer.mozilla.org/"),
        ("OWASP Foundation", "OWASP Top Ten Web Application Security Risks: Comprehensive Guidance for Defensive Web Architecture", "OWASP Project, 2021", "Available: https://owasp.org/www-project-top-ten/"),
        ("The Bootstrap Authors", "Bootstrap Icons Documentation (v1.11.3): Vector Iconography for Web Applications", "2024", "Available: https://icons.getbootstrap.com/"),
        ("David Shim", "QRCode.js: Cross-browser JavaScript QR Code Generator on HTML5 Canvas", "2020", "Available: https://github.com/davidshimjs/qrcodejs")
    ]
    
    for idx, (auth, title, pub, url) in enumerate(references_list, 1):
        p_ref = doc.add_paragraph()
        p_ref.paragraph_format.space_before = Pt(0)
        p_ref.paragraph_format.space_after = Pt(6)
        p_ref.paragraph_format.line_spacing = 1.15
        p_ref.paragraph_format.left_indent = Inches(0.35)
        p_ref.paragraph_format.first_line_indent = Inches(-0.35)
        
        r_num = p_ref.add_run(f"[{idx}] ")
        r_num.font.name = "Calibri"
        r_num.font.bold = True
        r_num.font.color.rgb = COLOR_HEADING
        
        r_body = p_ref.add_run(f"{auth}, \"{title},\" {pub}. {url}")
        r_body.font.name = "Calibri"
        r_body.font.color.rgb = COLOR_HEADING
        
    # =========================================================================
    # SAVE DESTINATIONS
    # =========================================================================
    output_paths = [
        os.path.join(BASE_DIR, "Healthy_Bite_Short_and_Simple_Documentation.docx"),
        os.path.join(BASE_DIR, "Documentation", "Healthy_Bite_Short_and_Simple_Documentation.docx"),
        os.path.join(BASE_DIR, "Project Doc Demo.docx") # Directly update user's working demo file!
    ]
    
    # Backup original Project Doc Demo.docx first
    demo_file = os.path.join(BASE_DIR, "Project Doc Demo.docx")
    backup_file = os.path.join(BASE_DIR, "Documentation", "Project Doc Demo_backup.docx")
    if os.path.exists(demo_file):
        shutil.copy2(demo_file, backup_file)
        print(f"Backed up original Project Doc Demo.docx to {backup_file}")
        
    for out_p in output_paths:
        try:
            doc.save(out_p)
            print(f"Successfully saved clean documentation to: {out_p}")
        except PermissionError:
            alt_p = out_p.replace(".docx", " (New Clean Version).docx")
            doc.save(alt_p)
            print(f"File was locked by Word: Saved clean copy to {alt_p} instead.")

if __name__ == "__main__":
    build_simple_documentation()
