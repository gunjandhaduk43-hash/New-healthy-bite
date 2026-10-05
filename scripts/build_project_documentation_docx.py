import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# --- Color Constants ---
COLOR_PRIMARY_HEX = "1B5E20"     # Forest Green (Healthy theme)
COLOR_SECONDARY_HEX = "2E7D32"   # Medium Green
COLOR_DARK_TEXT_HEX = "212121"   # Dark Charcoal
COLOR_MUTED_HEX = "555555"       # Slate Muted Text
COLOR_LIGHT_BG_HEX = "F4F6F4"    # Light Greenish Gray Tint
COLOR_BORDER_HEX = "D0D7DE"      # Border Gray
COLOR_PASS_HEX = "E8F5E9"        # Pass Light Green Fill
COLOR_PASS_TEXT_HEX = "1B5E20"   # Pass Dark Green Text

COLOR_PRIMARY = RGBColor(0x1B, 0x5E, 0x20)
COLOR_SECONDARY = RGBColor(0x2E, 0x7D, 0x32)
COLOR_DARK_TEXT = RGBColor(0x21, 0x21, 0x21)
COLOR_MUTED = RGBColor(0x55, 0x55, 0x55)

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for b_name, b_val in borders.items():
        if b_val:
            node = OxmlElement(f'w:{b_name}')
            node.set(qn('w:val'), b_val.get('val', 'single'))
            node.set(qn('w:sz'), str(b_val.get('sz', 4)))
            node.set(qn('w:space'), '0')
            node.set(qn('w:color'), b_val.get('color', COLOR_BORDER_HEX))
            tcBorders.append(node)
        else:
            node = OxmlElement(f'w:{b_name}')
            node.set(qn('w:val'), 'none')
            tcBorders.append(node)
    tcPr.append(tcBorders)

def add_header_footer(doc):
    for s in doc.sections:
        s.header_distance = Inches(0.5)
        s.footer_distance = Inches(0.5)
        
        # Header
        header = s.header
        p_head = header.paragraphs[0]
        p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_head = p_head.add_run("Healthy Bite — Digital Restaurant Menu & Food Ordering System")
        r_head.font.name = "Calibri"
        r_head.font.size = Pt(8.5)
        r_head.font.italic = True
        r_head.font.color.rgb = COLOR_MUTED
        
        # Footer
        footer = s.footer
        p_foot = footer.paragraphs[0]
        p_foot.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r_foot_left = p_foot.add_run("BCA Final Year Project Report • Department of Computer Applications, Christ College, Rajkot")
        r_foot_left.font.name = "Calibri"
        r_foot_left.font.size = Pt(8.5)
        r_foot_left.font.color.rgb = COLOR_MUTED

def style_heading_1(p, text):
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY
    return p

def style_heading_2(p, text):
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = COLOR_SECONDARY
    return p

def style_heading_3(p, text):
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = COLOR_DARK_TEXT
    return p

def add_body_paragraph(doc, text="", space_after=6, line_spacing=1.15):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if text:
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(10.5)
        run.font.color.rgb = COLOR_DARK_TEXT
    return p

def add_bullet_point(doc, bold_prefix, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    r_bold = p.add_run(bold_prefix)
    r_bold.font.name = "Calibri"
    r_bold.font.size = Pt(10.5)
    r_bold.font.bold = True
    r_bold.font.color.rgb = COLOR_DARK_TEXT
    
    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(10.5)
    r_text.font.color.rgb = COLOR_DARK_TEXT
    return p

def add_table_styled(doc, headers, data, col_widths=None):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    border_spec = {'val': 'single', 'sz': 4, 'color': COLOR_BORDER_HEX}
    
    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], COLOR_PRIMARY_HEX)
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)
        set_cell_borders(hdr_cells[i], top=border_spec, bottom=border_spec, left=border_spec, right=border_spec)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            
    # Data Rows
    for row_idx, row_data in enumerate(data):
        row_cells = table.rows[row_idx + 1].cells
        bg_color = COLOR_LIGHT_BG_HEX if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(row_data):
            row_cells[col_idx].text = str(val)
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], top=80, bottom=80, left=140, right=140)
            set_cell_borders(row_cells[col_idx], top=border_spec, bottom=border_spec, left=border_spec, right=border_spec)
            p = row_cells[col_idx].paragraphs[0]
            # Alignment logic
            if col_idx == 0 and ("TC-" in str(val) or str(val).isdigit()):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif "PASS" in str(val):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_cell_background(row_cells[col_idx], COLOR_PASS_HEX)
            elif any(k in headers[col_idx].lower() for k in ["status", "count", "type", "key"]):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(9.5)
                if "PASS" in str(val):
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0x1B, 0x5E, 0x20)
                elif col_idx == 0 and ("TC-" in str(val) or "TABLE:" in str(val)):
                    r.font.bold = True
                    r.font.color.rgb = COLOR_DARK_TEXT
                else:
                    r.font.color.rgb = COLOR_DARK_TEXT
                    
    # Column Widths
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = width
                
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table

def add_image_figure(doc, image_path, fig_num, caption, width=Inches(6.0)):
    if not os.path.exists(image_path):
        p_err = doc.add_paragraph()
        p_err.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p_err.add_run(f"[Image Missing: {image_path}]")
        r.font.italic = True
        r.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)
        return
        
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(4)
    run_img = p_img.add_run()
    run_img.add_picture(image_path, width=width)
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_after = Pt(12)
    p_cap.paragraph_format.keep_with_next = False
    
    r_fig = p_cap.add_run(f"Figure {fig_num}: ")
    r_fig.font.name = "Calibri"
    r_fig.font.size = Pt(9.5)
    r_fig.font.bold = True
    r_fig.font.color.rgb = COLOR_PRIMARY
    
    r_text = p_cap.add_run(caption)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(9.5)
    r_text.font.italic = True
    r_text.font.color.rgb = COLOR_MUTED

def add_code_snippet(doc, filename, module, purpose, line_ref, code_snippet, explanation, expected_behavior):
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(10)
    p_meta.paragraph_format.space_after = Pt(2)
    p_meta.paragraph_format.keep_with_next = True
    
    r_fn = p_meta.add_run(f"Source File: {filename} ")
    r_fn.font.name = "Calibri"
    r_fn.font.size = Pt(10)
    r_fn.font.bold = True
    r_fn.font.color.rgb = COLOR_PRIMARY
    
    r_mod = p_meta.add_run(f"| Module: {module} | Reference: {line_ref}\n")
    r_mod.font.name = "Calibri"
    r_mod.font.size = Pt(9.5)
    r_mod.font.italic = True
    r_mod.font.color.rgb = COLOR_MUTED
    
    r_purp = p_meta.add_run(f"Purpose: {purpose}")
    r_purp.font.name = "Calibri"
    r_purp.font.size = Pt(9.5)
    r_purp.font.bold = True
    r_purp.font.color.rgb = COLOR_DARK_TEXT
    
    # Code Table Box
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8F9FA")
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    set_cell_borders(cell, 
                     top={'val': 'single', 'sz': 4, 'color': 'E1E4E8'},
                     bottom={'val': 'single', 'sz': 4, 'color': 'E1E4E8'},
                     left={'val': 'single', 'sz': 16, 'color': COLOR_PRIMARY_HEX},
                     right={'val': 'single', 'sz': 4, 'color': 'E1E4E8'})
    
    p_code = cell.paragraphs[0]
    p_code.paragraph_format.space_after = Pt(0)
    p_code.paragraph_format.line_spacing = 1.05
    r_c = p_code.add_run(code_snippet.strip())
    r_c.font.name = "Consolas"
    r_c.font.size = Pt(8.5)
    r_c.font.color.rgb = RGBColor(0x24, 0x29, 0x2E)
    
    p_exp = doc.add_paragraph()
    p_exp.paragraph_format.space_before = Pt(4)
    p_exp.paragraph_format.space_after = Pt(2)
    p_exp.paragraph_format.line_spacing = 1.15
    r_exp_b = p_exp.add_run("Technical Explanation: ")
    r_exp_b.font.name = "Calibri"
    r_exp_b.font.size = Pt(10)
    r_exp_b.font.bold = True
    r_exp_t = p_exp.add_run(explanation)
    r_exp_t.font.name = "Calibri"
    r_exp_t.font.size = Pt(10)
    
    p_beh = doc.add_paragraph()
    p_beh.paragraph_format.space_after = Pt(12)
    p_beh.paragraph_format.line_spacing = 1.15
    r_beh_b = p_beh.add_run("Expected System Behavior: ")
    r_beh_b.font.name = "Calibri"
    r_beh_b.font.size = Pt(10)
    r_beh_b.font.bold = True
    r_beh_b.font.color.rgb = COLOR_SECONDARY
    r_beh_t = p_beh.add_run(expected_behavior)
    r_beh_t.font.name = "Calibri"
    r_beh_t.font.size = Pt(10)

print("Base styling & helper functions defined successfully.")
